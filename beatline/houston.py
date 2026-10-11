"""Houston, TX public-data connectors.

* Beats:  City of Houston GIS (COHGIS) Public_Safety map service, "Houston
          Police Beats" layer. COHGIS data is public domain.
* Calls:  HEC Active Incidents page (police), which the city regenerates every
          five minutes from its dispatch system. It lists address, cross
          street, Key Map grid, call time and incident type - no coordinates
          and no priority - so addresses are geocoded with the free US Census
          geocoder and priority is recorded as 0 (unknown).
* Baseline: HPD's monthly crime spreadsheet (NIBRS) by street and beat, used as
          "where calls usually come from" for each beat.

There is no public source for officer locations; the dashboard runs in its
calls-only mode until an agency connects its own AVL feed.
"""

from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
import sys
import time
from datetime import datetime
from html.parser import HTMLParser
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from zoneinfo import ZoneInfo

from . import analytics, db, importers

GIS_SERVICE = "https://mycity2.houstontx.gov/pubgis02/rest/services/HoustonMap/Public_Safety/MapServer"
BEATS_LAYER_DEFAULT = 5          # "Houston Police Beats"; looked up by name first
ACTIVE_INCIDENTS_URL = "https://cohweb.houstontx.gov/ActiveIncidents/Combined.aspx?agency=P"
CENSUS_GEOCODER = "https://geocoding.geo.census.gov/geocoder/locations/onelineaddress"
TIMEZONE = "America/Chicago"
USER_AGENT = "Beatline/0.1 (patrol coverage research; +https://github.com/AmitBhalesha/AmitBhalesha)"
POLL_SECONDS = 300               # the source updates every five minutes; polling faster gains nothing


def http_get(url: str, params: dict | None = None, timeout: float = 30) -> bytes:
    if params:
        url += ("&" if "?" in url else "?") + urlencode(params)
    with urlopen(Request(url, headers={"User-Agent": USER_AGENT}), timeout=timeout) as resp:
        return resp.read()


def _norm(text: str) -> str:
    return re.sub(r"[^a-z0-9#]", "", str(text).lower()).replace("#", "")


# ---------------------------------------------------------------- beats (GIS)

def find_beats_layer(fetch=http_get) -> int:
    info = json.loads(fetch(GIS_SERVICE, {"f": "json"}))
    for layer in info.get("layers", []):
        if "beat" in layer.get("name", "").lower() and "police" in layer.get("name", "").lower():
            return layer["id"]
    return BEATS_LAYER_DEFAULT


def fetch_beats_geojson(fetch=http_get, layer: int | None = None) -> dict:
    layer = find_beats_layer(fetch) if layer is None else layer
    features: list = []
    while True:
        page = json.loads(fetch(f"{GIS_SERVICE}/{layer}/query", {
            "where": "1=1", "outFields": "*", "returnGeometry": "true", "outSR": 4326, "f": "geojson",
            "resultOffset": len(features), "resultRecordCount": 1000,
        }))
        if "error" in page:
            raise RuntimeError(f"GIS query failed: {page['error']}")
        batch = page.get("features", [])
        features += batch
        more = page.get("exceededTransferLimit") or (page.get("properties") or {}).get("exceededTransferLimit")
        if not batch or not more:
            break
    return {"type": "FeatureCollection", "features": features}


def setup(conn: sqlite3.Connection, fetch=http_get, layer: int | None = None) -> dict:
    """Load real HPD beats and configure the database for Houston public data."""
    db.init_db(conn)
    fc = fetch_beats_geojson(fetch, layer)
    count = importers.import_beats_geojson(conn, fc)
    if count == 0:
        raise RuntimeError("The GIS service returned no beat polygons")
    db.set_meta(conn, "agency_name", "Houston, TX (public data)")
    db.set_meta(conn, "timezone", TIMEZONE)
    db.set_meta(conn, "demand_source", "crime")
    db.set_meta(conn, "data_note", "Public data · no officer GPS")
    conn.commit()
    sample = fc["features"][0].get("properties", {}) if fc["features"] else {}
    return {"beats": count, "fields": sorted(sample)}


# ------------------------------------------------------ live calls (HEC page)

class _Tables(HTMLParser):
    """Collects every <table> as a list of rows of cell text."""

    def __init__(self):
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self._stack: list[list[list[str]]] = []
        self._row: list[str] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        if tag == "table":
            self._stack.append([])
        elif tag == "tr" and self._stack:
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append(re.sub(r"\s+", " ", "".join(self._cell)).strip())
            self._cell = None
        elif tag == "tr" and self._row is not None and self._stack:
            if self._row:
                self._stack[-1].append(self._row)
            self._row = None
        elif tag == "table" and self._stack:
            self.tables.append(self._stack.pop())

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


COLUMN_KEYS = {
    "agency": ("agency",),
    "address": ("address", "location"),
    "cross_street": ("crossstreet", "cross"),
    "key_map": ("keymap",),
    "call_time": ("calltimeopened", "calltime", "callopened", "opened", "time"),
    "incident_type": ("incidenttype", "type", "incident"),
    "combined": ("combinedresponse", "combined"),
}
TIME_FORMATS = ("%m/%d/%Y %H:%M", "%m/%d/%Y %H:%M:%S", "%m/%d/%Y %I:%M %p", "%m/%d/%Y %I:%M:%S %p",
                "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M")


def _map_header(cells: list[str]) -> dict[str, int]:
    found: dict[str, int] = {}
    normed = [_norm(c) for c in cells]
    for key, names in COLUMN_KEYS.items():
        for name in names:
            idx = next((i for i, n in enumerate(normed) if n == name and i not in found.values()), None)
            if idx is not None:
                found[key] = idx
                break
    return found


def parse_time(text: str, tz: ZoneInfo) -> int | None:
    text = re.sub(r"\s+", " ", text or "").strip()
    for fmt in TIME_FORMATS:
        try:
            return int(datetime.strptime(text, fmt).replace(tzinfo=tz).timestamp())
        except ValueError:
            continue
    return None


def parse_active_incidents(html: str, tz_name: str = TIMEZONE) -> tuple[list[dict], dict]:
    """Returns (police incidents, diagnostics). Columns are matched by header
    name, so reordered or renamed-in-small-ways columns still parse."""
    tz = ZoneInfo(tz_name)
    parser = _Tables()
    parser.feed(html)
    diag = {"tables": len(parser.tables), "header": None, "rows": 0, "skipped": 0}
    for rows in parser.tables:
        for h, header in enumerate(rows[:5]):
            cols = _map_header(header)
            if "address" in cols and "call_time" in cols:
                break
        else:
            continue
        diag["header"] = header
        out = []
        for row in rows[h + 1:]:
            if len(row) < len(header):
                continue
            diag["rows"] += 1
            get = lambda key: row[cols[key]] if key in cols else ""  # noqa: E731
            agency = get("agency")
            if agency and not re.match(r"^(p|hpd|police)", agency.strip(), re.I):
                continue
            ts = parse_time(get("call_time"), tz)
            if not get("address") or ts is None:
                diag["skipped"] += 1
                continue
            key = "|".join([get("address").upper(), get("cross_street").upper(), get("call_time")])
            out.append({
                "id": "HEC-" + hashlib.sha1(key.encode()).hexdigest()[:12],
                "ts": ts,
                "address": get("address"),
                "cross_street": get("cross_street"),
                "key_map": get("key_map"),
                "incident_type": get("incident_type") or "Unknown",
            })
        return out, diag
    return [], diag


def clean_address(address: str) -> str:
    a = re.sub(r"\s+", " ", address.upper()).strip()
    a = re.sub(r"^(\d+)\s*-\s*\d+\b", r"\1", a)               # "1200-1299 MAIN" -> "1200 MAIN"
    a = re.sub(r"\b(BLOCK OF|BLK OF|BLOCK|BLK)\b", " ", a)
    return re.sub(r"\s+", " ", a).strip(" ,")


def geocode(conn: sqlite3.Connection, address: str, fetch=http_get) -> tuple[float, float] | None:
    """Census geocoder with a permanent cache (misses are cached too)."""
    key = clean_address(address)
    row = conn.execute("SELECT lat, lon FROM geocode_cache WHERE address = ?", (key,)).fetchone()
    if row:
        return (row["lat"], row["lon"]) if row["lat"] is not None else None
    resp = json.loads(fetch(CENSUS_GEOCODER, {
        "address": f"{key}, HOUSTON, TX", "benchmark": "Public_AR_Current", "format": "json",
    }))
    matches = (resp.get("result") or {}).get("addressMatches") or []
    lat = lon = matched = None
    if matches:
        lon, lat = matches[0]["coordinates"]["x"], matches[0]["coordinates"]["y"]
        matched = matches[0].get("matchedAddress")
    conn.execute(
        "INSERT OR REPLACE INTO geocode_cache (address, lat, lon, matched, ts) VALUES (?, ?, ?, ?, ?)",
        (key, lat, lon, matched, int(time.time())),
    )
    return (lat, lon) if lat is not None else None


def poll_once(conn: sqlite3.Connection, fetch=http_get, now: int | None = None) -> dict:
    """Fetch the Active Incidents page once and record new/ongoing police calls.
    A call's cleared_ts is the last time it was still listed."""
    now = now or int(time.time())
    html = fetch(ACTIVE_INCIDENTS_URL).decode("utf-8", "replace")
    incidents, diag = parse_active_incidents(html)
    stats = {"listed": len(incidents), "new": 0, "ongoing": 0, "unmapped": 0, "geocode_errors": 0}
    if diag["header"] is None:
        stats["warning"] = "No incidents table found; the page layout may have changed (run `houston check`)."
    for inc in incidents:
        if conn.execute("UPDATE calls SET cleared_ts = ? WHERE id = ?", (now, inc["id"])).rowcount:
            stats["ongoing"] += 1
            continue
        if conn.execute("UPDATE unmapped_calls SET last_seen = ? WHERE id = ?", (now, inc["id"])).rowcount:
            stats["ongoing"] += 1
            continue
        try:
            loc = geocode(conn, inc["address"], fetch)
        except Exception:  # network hiccup: leave it for the next poll instead of marking it unmapped
            stats["geocode_errors"] += 1
            continue
        if loc is None:
            conn.execute(
                "INSERT OR IGNORE INTO unmapped_calls (id, ts, call_type, address, last_seen) VALUES (?, ?, ?, ?, ?)",
                (inc["id"], inc["ts"], inc["incident_type"], inc["address"], now),
            )
            stats["unmapped"] += 1
            continue
        conn.execute(
            "INSERT OR IGNORE INTO calls (id, ts, cleared_ts, call_type, priority, lat, lon, officer_id, source, address) "
            "VALUES (?, ?, ?, ?, 0, ?, ?, NULL, 'hec', ?)",
            (inc["id"], inc["ts"], now, inc["incident_type"], loc[0], loc[1],
             inc["address"] + (f" @ {inc['cross_street']}" if inc["cross_street"] else "")),
        )
        stats["new"] += 1
    conn.commit()
    return stats


def poll_forever(conn: sqlite3.Connection, every: int = POLL_SECONDS, fetch=http_get) -> None:
    every = max(every, 60)
    while True:
        started = time.time()
        stamp = datetime.now(ZoneInfo(TIMEZONE)).strftime("%Y-%m-%d %H:%M:%S")
        try:
            stats = poll_once(conn, fetch)
            print(f"{stamp} {json.dumps(stats)}", flush=True)
        except Exception as exc:  # keep polling through outages
            print(f"{stamp} poll failed: {exc}", file=sys.stderr, flush=True)
        time.sleep(max(5, every - (time.time() - started)))


# --------------------------------------------------- crime baseline (monthly)

CRIME_COLUMNS = {
    "incident": ("incident", "incidentnumber", "incidentno"),
    "date": ("rmsoccurrencedate", "occurrencedate", "date", "offensedate"),
    "hour": ("rmsoccurrencehour", "occurrencehour", "hour"),
    "offense": ("nibrsdescription", "offensetype", "offense", "nibrsclass"),
    "count": ("offensecount", "ofoffenses", "numoffenses", "offenses"),
    "beat": ("beat",),
    "lat": ("maplatitude", "latitude", "lat"),
    "lon": ("maplongitude", "longitude", "lon", "long"),
    "block": ("blockrange", "streetno", "block"),
    "street": ("streetname", "street"),
    "street_type": ("streettype",),
}


def _rows(path: str):
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            yield from csv.reader(f)
        return
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("Reading .xlsx needs openpyxl: pip install -e '.[houston]'") from exc
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            yield list(row)


def _crime_header(row: list) -> dict[str, int] | None:
    normed = [_norm(c) if c is not None else "" for c in row]
    cols = {}
    for key, names in CRIME_COLUMNS.items():
        for name in names:
            if name in normed:
                cols[key] = normed.index(name)
                break
    return cols if "date" in cols and "offense" in cols and ("beat" in cols or "lat" in cols) else None


def _to_date(value, tz: ZoneInfo) -> datetime | None:
    if isinstance(value, datetime):
        return value.replace(tzinfo=tz) if value.tzinfo is None else value
    if hasattr(value, "year") and hasattr(value, "month"):
        return datetime(value.year, value.month, value.day, tzinfo=tz)
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S", "%m/%d/%Y %I:%M:%S %p", "%m/%d/%Y %H:%M"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=tz)
        except ValueError:
            continue
    return None


def _float(value) -> float | None:
    try:
        f = float(value)
    except (TypeError, ValueError):
        return None
    return f if f == f and abs(f) > 1 else None  # 0 / blank / NaN mean "no coordinate"


def import_crime(conn: sqlite3.Connection, path: str) -> dict:
    """Load an HPD monthly crime file (.xlsx or .csv). Columns are found by
    header name; a row needs a date, an offense, and a beat or coordinates."""
    tz = ZoneInfo(db.get_meta(conn, "timezone", TIMEZONE))
    beats = set(db.load_beats(conn))
    cols = None
    stats = {"rows": 0, "imported": 0, "skipped": 0, "beats_matched": 0, "with_coordinates": 0,
             "first": None, "last": None}
    batch = []
    for row in _rows(path):
        if cols is None:
            cols = _crime_header(row)
            continue
        if not any(v not in (None, "") for v in row):
            continue
        stats["rows"] += 1
        get = lambda key: row[cols[key]] if key in cols and cols[key] < len(row) else None  # noqa: E731
        day = _to_date(get("date"), tz)
        offense = str(get("offense") or "").strip()
        if day is None or not offense:
            stats["skipped"] += 1
            continue
        hour = get("hour")
        try:
            day = day.replace(hour=int(float(hour))) if hour not in (None, "") and day.hour == 0 else day
        except (ValueError, TypeError):
            pass
        beat = str(get("beat") or "").strip().upper() or None
        lat, lon = _float(get("lat")), _float(get("lon"))
        if beat not in beats and lat is None:
            stats["skipped"] += 1
            continue
        try:
            count = max(1, int(float(get("count") or 1)))
        except (TypeError, ValueError):
            count = 1
        block = " ".join(str(get(k)).strip() for k in ("block", "street", "street_type") if get(k) not in (None, ""))
        ts = int(day.timestamp())
        row_id = hashlib.sha1("|".join("" if v is None else str(v) for v in row).encode()).hexdigest()[:20]
        batch.append((row_id, ts, offense, count, beat if beat in beats else None, lat, lon, block or None))
        stats["imported"] += 1
        stats["beats_matched"] += beat in beats
        stats["with_coordinates"] += lat is not None
        stats["first"] = ts if stats["first"] is None else min(stats["first"], ts)
        stats["last"] = ts if stats["last"] is None else max(stats["last"], ts)
        if len(batch) >= 5000:
            _insert_crime(conn, batch)
            batch = []
    if cols is None:
        raise ValueError(f"No header row with a date, an offense and a beat/coordinates column in {path}")
    _insert_crime(conn, batch)
    db.set_meta(conn, "demand_source", "crime")
    conn.commit()
    for k in ("first", "last"):
        if stats[k] is not None:
            stats[k] = datetime.fromtimestamp(stats[k], tz).date().isoformat()
    return stats


def _insert_crime(conn: sqlite3.Connection, rows: list[tuple]) -> None:
    conn.executemany(
        "INSERT OR IGNORE INTO crime_reports (id, ts, offense, offense_count, beat_id, lat, lon, block_address) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        rows,
    )


# ----------------------------------------------------------------- diagnostics

def check(conn: sqlite3.Connection, fetch=http_get) -> None:
    """Hit every live source once and print what Beatline sees - run this first
    if anything looks wrong, and paste the output when asking for help."""
    print("1. Beats (COHGIS)")
    try:
        layer = find_beats_layer(fetch)
        fc = fetch_beats_geojson(fetch, layer)
        props = fc["features"][0]["properties"] if fc["features"] else {}
        print(f"   layer {layer}: {len(fc['features'])} features; fields: {sorted(props)}")
        print(f"   beat id field: {importers.beat_id_field(props) if props else '?'}; example: {props}")
    except Exception as exc:
        print(f"   FAILED: {exc}")
    print("2. Active Incidents (HEC)")
    sample = None
    try:
        incidents, diag = parse_active_incidents(fetch(ACTIVE_INCIDENTS_URL).decode("utf-8", "replace"))
        print(f"   {diag['tables']} tables; header: {diag['header']}; rows: {diag['rows']}; police incidents: {len(incidents)}")
        for inc in incidents[:3]:
            print(f"   {inc}")
        sample = incidents[0]["address"] if incidents else None
    except Exception as exc:
        print(f"   FAILED: {exc}")
    print("3. Geocoder (US Census)")
    try:
        address = sample or "901 BAGBY ST"
        print(f"   {address!r} -> {geocode(conn, address, fetch)}")
        conn.rollback()
    except Exception as exc:
        print(f"   FAILED: {exc}")
    beats = len(db.load_beats(conn))
    crimes = conn.execute("SELECT COUNT(*), MIN(ts), MAX(ts) FROM crime_reports").fetchone()
    calls = conn.execute("SELECT COUNT(*) FROM calls WHERE source = 'hec'").fetchone()[0]
    unmapped = conn.execute("SELECT COUNT(*) FROM unmapped_calls").fetchone()[0]
    print(f"4. Database: {beats} beats, {calls} live calls ({unmapped} unmapped), {crimes[0]} crime rows")
    print(f"   demand baseline: {analytics.demand_points(conn, int(time.time()))[1]}")
